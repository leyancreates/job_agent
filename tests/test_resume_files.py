from io import BytesIO

import pytest
from docx import Document

from resumes import (
    CorruptedResumeFileError,
    EmptyResumeError,
    UnsupportedResumeFileError,
    resume_from_google_doc,
    resume_from_upload,
)


def minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    content = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{number} 0 obj\n".encode("ascii"))
        content.extend(body)
        content.extend(b"\nendobj\n")
    xref_offset = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    content.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    content.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(content)


def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Jane Doe", level=1)
    document.add_paragraph("Python and SQL analyst")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Toronto"
    table.cell(0, 1).text = "Data Analytics"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_pdf_parsing_extracts_plain_text():
    resume = resume_from_upload("resume.PDF", minimal_pdf("Jane Doe Python SQL"))
    assert resume.source == "Upload File"
    assert resume.filename == "resume.PDF"
    assert "Jane Doe Python SQL" in resume.text
    assert resume.word_count == 4


def test_docx_parsing_extracts_paragraphs_and_tables():
    resume = resume_from_upload("resume.docx", docx_bytes())
    assert "Jane Doe" in resume.text
    assert "Python and SQL analyst" in resume.text
    assert "Toronto | Data Analytics" in resume.text


def test_txt_parsing_accepts_utf8_and_bom():
    resume = resume_from_upload(
        "resume.txt", "\ufeffJane Doe\r\nPython developer".encode("utf-8")
    )
    assert resume.text == "Jane Doe\nPython developer"
    assert resume.word_count == 4


@pytest.mark.parametrize(
    ("filename", "data"),
    [
        ("empty.txt", b""),
        ("whitespace.txt", b" \n\t"),
    ],
)
def test_empty_files_are_rejected(filename, data):
    with pytest.raises(EmptyResumeError, match="empty|no readable text"):
        resume_from_upload(filename, data)


@pytest.mark.parametrize(
    ("filename", "data", "message"),
    [
        ("broken.pdf", b"not a pdf", "Could not read this PDF"),
        ("broken.docx", b"not a docx", "Could not read this DOCX"),
        ("broken.txt", b"\xff\xfe\x00", "Save it as UTF-8"),
    ],
)
def test_corrupted_files_show_format_specific_errors(filename, data, message):
    with pytest.raises(CorruptedResumeFileError, match=message):
        resume_from_upload(filename, data)


def test_unsupported_file_is_rejected():
    with pytest.raises(UnsupportedResumeFileError, match="PDF, DOCX, or UTF-8 TXT"):
        resume_from_upload("resume.rtf", b"resume")


def test_google_docs_uses_the_same_resume_model():
    resume = resume_from_google_doc("Jane Doe\nData Analyst", "doc-123")
    assert resume.source == "Google Docs"
    assert resume.google_doc_id == "doc-123"
    assert resume.word_count == 4
