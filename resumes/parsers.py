"""Extract normalized plain text from supported resume sources."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pdfplumber
from docx import Document

from .errors import (
    CorruptedResumeFileError,
    EmptyResumeError,
    UnsupportedResumeFileError,
)
from .models import ResumeContent

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def _normalize_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    return "\n".join(line.rstrip() for line in normalized.splitlines()).strip()


def _validated_resume(
    text: str,
    *,
    source: str,
    filename: str,
    google_doc_id: str | None = None,
) -> ResumeContent:
    normalized = _normalize_text(text)
    if not normalized:
        raise EmptyResumeError("The resume contains no readable text.")
    return ResumeContent(
        text=normalized,
        source=source,
        filename=filename,
        google_doc_id=google_doc_id,
    )


def _extract_pdf(data: bytes) -> str:
    try:
        with pdfplumber.open(BytesIO(data)) as pdf:
            return "\n\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as exc:
        raise CorruptedResumeFileError(
            "Could not read this PDF. It may be corrupted or password-protected."
        ) from exc


def _extract_docx(data: bytes) -> str:
    try:
        document = Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    parts.append(row_text)
        return "\n".join(parts)
    except Exception as exc:
        raise CorruptedResumeFileError(
            "Could not read this DOCX file. It may be corrupted or not a valid Word document."
        ) from exc


def _extract_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise CorruptedResumeFileError(
            "Could not read this TXT file. Save it as UTF-8 and upload it again."
        ) from exc


def resume_from_upload(filename: str, data: bytes) -> ResumeContent:
    extension = Path(filename).suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedResumeFileError(
            "Unsupported file type. Upload a PDF, DOCX, or UTF-8 TXT file."
        )
    if not data:
        raise EmptyResumeError("The uploaded resume is empty.")

    extractors = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".txt": _extract_txt,
    }
    return _validated_resume(
        extractors[extension](data),
        source="Upload File",
        filename=Path(filename).name,
    )


def resume_from_google_doc(text: str, doc_id: str) -> ResumeContent:
    return _validated_resume(
        text,
        source="Google Docs",
        filename="Google Docs resume",
        google_doc_id=doc_id,
    )
